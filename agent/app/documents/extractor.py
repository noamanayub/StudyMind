from dataclasses import dataclass
from pathlib import Path
import re
import zipfile
from pypdf import PdfReader
from docx import Document
from app.documents.ocr import pdf_page_text, image_text, OCRUnavailable
from app.documents.media import visual_analysis, transcribe_audio, transcribe_youtube
from app.llm.client import ProviderUnavailable


class InvalidDocument(ValueError):
    pass


@dataclass
class Section:
    text: str
    page: int | None = None
    title: str | None = None


def clean(text: str) -> str:
    return re.sub(r'[ \t]+', ' ', text.replace('\x00', '')).strip()


def extract(path: Path, mode: str = 'TEXT', heartbeat=None) -> list[Section]:
    ext = path.suffix.lower()
    try:
        if ext == '.pdf':
            reader = PdfReader(path)
            if reader.is_encrypted:
                raise InvalidDocument('Password-protected PDFs are not supported. Upload an unlocked copy.')
            if len(reader.pages) > 1500:
                raise InvalidDocument('This PDF is too large to process. Split it into smaller documents.')
            sections = [Section(clean(page.extract_text() or ''), i + 1) for i, page in enumerate(reader.pages)]
            missing = [s.page for s in sections if not s.text and reader.pages[s.page - 1].images]
            if missing:
                recognized = pdf_page_text(path, missing, heartbeat)
                for section in sections:
                    if not section.text:
                        section.text = clean(recognized.get(section.page, ''))
        elif ext in ('.png', '.jpg', '.jpeg', '.webp'):
            if mode == 'VISION':
                sections = [Section(clean(visual_analysis(path)), title='AI image analysis — verify against original')]
            else:
                sections = [Section(clean(image_text(path)), title='Image text (OCR)')]
        elif ext in ('.mp3', '.wav', '.m4a'):
            segments = transcribe_audio(path)
            sections = [Section(clean(s.text), title=f'AI transcript · {s.startSeconds // 60:02}:{s.startSeconds % 60:02}–{s.endSeconds // 60:02}:{s.endSeconds % 60:02}') for s in segments]
        elif ext == '.youtube':
            segments = transcribe_youtube(path)
            sections = [Section(clean(s.text), title=f'AI YouTube transcript · {s.startSeconds // 60:02}:{s.startSeconds % 60:02}–{s.endSeconds // 60:02}:{s.endSeconds % 60:02}') for s in segments]
        elif ext == '.docx':
            with zipfile.ZipFile(path) as archive:
                entries = archive.infolist()
                if len(entries) > 5000 or sum(e.file_size for e in entries) > 100 * 1024 * 1024:
                    raise InvalidDocument('This document expands beyond the processing limit.')
            doc = Document(path)
            sections = []
            for i, p in enumerate(doc.paragraphs):
                if p.text.strip():
                    sections.append(Section(clean(p.text), title=f'Paragraph {i + 1}'))
            for i, table in enumerate(doc.tables):
                sections.append(Section(clean('\n'.join(' | '.join(c.text for c in row.cells) for row in table.rows)), title=f'Table {i + 1}'))
        elif ext in ('.txt', '.md'):
            text = path.read_text(encoding='utf-8-sig')
            sections = [Section(clean(part), title=f'Section {i + 1}') for i, part in enumerate(re.split(r'\n\s*\n', text)) if part.strip()]
        else:
            raise InvalidDocument('This document type is not supported.')
    except OCRUnavailable as error:
        raise InvalidDocument(str(error)) from error
    except InvalidDocument:
        raise
    except ProviderUnavailable:
        raise
    except ValueError as error:
        raise InvalidDocument(str(error)) from error
    except Exception as error:
        raise InvalidDocument('We could not read this document. Check the file and upload it again.') from error
    sections = [s for s in sections if s.text]
    if not sections:
        raise InvalidDocument('No readable text was found. Use a clearer scan or image with legible text.')
    if sum(len(s.text) for s in sections) > 5_000_000:
        raise InvalidDocument('This document contains too much text. Split it into smaller files.')
    return sections
