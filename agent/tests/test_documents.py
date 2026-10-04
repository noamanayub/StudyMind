from pathlib import Path
import pytest
from docx import Document
from pypdf import PdfWriter
from app.documents.extractor import extract, InvalidDocument, Section
from app.rag.chunker import chunk_sections, encoder


def test_text_and_markdown_preserve_sections(tmp_path):
    for ext in ('.txt', '.md'):
        path = tmp_path / f'lecture{ext}'
        path.write_text('Normalization reduces duplication.\n\nA primary key identifies a row.', encoding='utf-8')
        sections = extract(path)
        assert len(sections) == 2
        assert sections[1].title == 'Section 2'
        assert sections[0].page is None


def test_docx_preserves_paragraph_and_table_locations(tmp_path):
    path = tmp_path / 'lecture.docx'
    doc = Document()
    doc.add_paragraph('A database stores organized information.')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Primary key'
    table.cell(0, 1).text = 'Unique identifier'
    doc.save(path)
    sections = extract(path)
    assert sections[0].title == 'Paragraph 1'
    assert sections[1].title == 'Table 1'
    assert 'Unique identifier' in sections[1].text


def test_invalid_empty_and_scanned_files_are_rejected(tmp_path):
    empty = tmp_path / 'empty.txt'
    empty.write_text('  \n ')
    with pytest.raises(InvalidDocument, match='No readable text'):
        extract(empty)
    fake = tmp_path / 'fake.pdf'
    fake.write_text('This is not a PDF')
    with pytest.raises(InvalidDocument, match='could not read'):
        extract(fake)
    scanned = tmp_path / 'scanned.pdf'
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(scanned)
    with pytest.raises(InvalidDocument, match='No readable text'):
        extract(scanned)


def test_chunks_preserve_page_metadata_and_token_bounds():
    sections = [Section('Database normalization. ' * 600, page=12)]
    chunks = chunk_sections(sections, 1000, 150)
    assert len(chunks) > 1
    assert all(c['page_number'] == 12 for c in chunks)
    assert all(len(encoder().encode(c['content'])) <= 1000 for c in chunks)
    assert [c['chunk_index'] for c in chunks] == list(range(len(chunks)))
    assert encoder().encode(chunks[0]['content'])[-150:] == encoder().encode(chunks[1]['content'])[:150]


def test_invalid_overlap():
    with pytest.raises(ValueError):
        chunk_sections([Section('Notes')], 100, 100)
